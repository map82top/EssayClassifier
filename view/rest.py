from flask import Flask, request, redirect, jsonify
from flask_socketio import SocketIO, emit, send
from pptx import Presentation
import pandas as pd
from analyzer.analyzer import Analyzer
from storage.core import run_orm
from storage.entities import ReportSchema
from downloader.downloader import download_archive
from downloader.gdrive import Drive
import traceback
import zipfile
from analyzer.exceptions import *
import json
import re
import docx
from analyzer.lecture_reader import read_from_presentation
import io

class EstimatorServer():
    def __init__(self, gdrive_certificat_path):
        self.server = Flask(__name__,
                       static_url_path='',
                       static_folder='frontend/static')

        self.socketio = SocketIO(self.server)
        self.session_maker = run_orm()
        self.analyzer = Analyzer()
        self.drive = Drive(gdrive_certificat_path)

        self.server.route('/', methods=['GET'])(self.index)
        self.server.route('/upload', methods=['POST'])(self.upload_task)
        self.server.route('/end_check', methods=['POST'])(self.end_check)

    def item_callback(self, current_item_id, count_items):
        self.socketio.emit('changed-report-status', json.dumps({'status': 'handling',
                            'description': f'Обработано ответов из архива {current_item_id} из {count_items}'}))

    def load_lecture_file(self, file):
        extension = re.findall(r'\.\w+$', file.filename)[0]

        if extension == '.pptx':
            lecture = Presentation(file)
            return read_from_presentation(lecture)
        if extension == '.txt':
            wrapper = io.TextIOWrapper(file, encoding='utf-8')
            return wrapper.read()
        if extension == '.docx':
            lecture = docx.Document(file)
            fullText = []
            for para in lecture.paragraphs:
                fullText.append(para.text)
            return ' '.join(fullText)

        raise NotSupportLectureExtensionType(extension)

    def load_essays_file(self, file):
        extension = re.findall(r'\.\w+$', file.filename)[0]

        if extension == '.csv':
            return pd.read_csv(file)
        if extension == '.xlsx':
            return pd.read_excel(file)
        if extension == '.zip':
            archive = zipfile.ZipFile(file, 'r')
            return download_archive(self.drive, archive, self.item_callback)

        raise NotSupportEssayExtensionType(extension)

    def index(self):
        return self.server.send_static_file("index.html")

    def upload_task(self):
        try:
            if 'lecture' not in request.files or 'essays' not in request.files:
                return redirect(request.url)

            lecture = request.files['lecture']
            essays = request.files['essays']

            if lecture.filename == '' or essays.filename == '':
                return redirect(request.url)

            self.socketio.emit('changed-report-status', json.dumps({'status': 'handling', "description": "Обработка файлов"}))

            lecture = self.load_lecture_file(lecture)
            essays = self.load_essays_file(essays)
            essays = essays.dropna(axis=0)

            self.socketio.emit('changed-report-status', json.dumps({'status': 'handling', "description": "Оценка эссе"}))

            report = self.analyzer.analyze(lecture, essays)
            report_schema = ReportSchema()
            session = self.session_maker()
            session.add(report)
            session.commit()

            self.socketio.emit('changed-report-status', json.dumps({'status': 'handled'}))
            return report_schema.dump(report)

        except Exception as e:
            print(e)
            print(traceback.print_exc())

            if type(e) == NotFoundEssayColumn or type(e) == NotSupportEssayExtensionType\
                    or type(e) == NotSupportLectureExtensionType:
                return json.dumps({"status": "error", "text": str(e)}), 500
            else:
                return json.dumps({"status": "error", "text": "Ошибка оценки загруженных эссе"}), 500


    def end_check(self):
        try:
            session = self.session_maker()
            report_schema = ReportSchema()
            json_data = json.loads(request.data)
            report = report_schema.load(data=json_data, session=session)

            if report.lecture is None:
                raise Exception("Received incorrect data. Report not found")

            session.commit()
            return json.dumps({"status": "success", "text": "Результаты проверки эссе успешно сохранены"})

        except Exception as e:
            print(e)
            print(traceback.print_exc())
            session.close()
            return json.dumps({"status": "error", "text": "Ошибка сохранения резултатов проверки"}), 500

    def start(self):
        self.socketio.run(self.server, host='localhost', port=5000)


if __name__ == '__main__':
    server = EstimatorServer("../downloader/resources")
    server.start()

